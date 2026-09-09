import os
import pandas as pd
import numpy as np
import openpyxl
import yfinance as yf
import pandas_ta as ta
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage

def generiraj_excel_izvjestaj(csv_path="podaci.csv", excel_path="Financijski_Izvjestaj.xlsx"):
    if not os.path.exists(csv_path):
        print(f"[GREŠKA] Datoteka {csv_path} ne postoji. Prvo pokreni Nabava_podataka.py!")
        return

    print("[POKRETANJE] Generiranje Excel izvještaja...")
    
    # Učitavanje podataka i rješavanje MultiIndex formata iz yfinance-a
    try:
        df = pd.read_csv(csv_path, header=[0, 1], index_col=0, parse_dates=True)
    except Exception:
        df = pd.read_csv(csv_path, index_col=0, parse_dates=True)
    
    # Izdvajanje cijena zatvaranja ('Close')
    if isinstance(df.columns, pd.MultiIndex):
        if 'Close' in df.columns.levels[0]:
            df_prices = df['Close'].ffill().bfill()
        else:
            df_prices = df.ffill().bfill()
    else:
        df_prices = df.ffill().bfill()

    tickers = list(df_prices.columns)
    df_returns = df_prices.pct_change().dropna()
    dates = df_prices.index

    # Inicijalizacija Excel radne knjige
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Ukloni zadani prazni list

    # Stilovi (Cool Tech Palette: Tamnoplava i Siva)
    HEADER_BG = "203764"      # Tamnoplava
    HEADER_FG = "FFFFFF"      # Bijela
    ACCENT_BG = "D9E1F2"      # Svijetlo plava (ledena)
    ZEBRA_BG = "F2F2F2"       # Svijetlo siva
    BORDER_COLOR = "D9D9D9"   # Siva za obrube

    font_title = Font(name="Arial", size=16, bold=True, color="203764")
    font_subtitle = Font(name="Arial", size=10, italic=True, color="595959")
    font_section = Font(name="Arial", size=12, bold=True, color="203764")
    font_header = Font(name="Arial", size=10, bold=True, color=HEADER_FG)
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=10)

    fill_header = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    fill_accent = PatternFill(start_color=ACCENT_BG, end_color=ACCENT_BG, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")

    thin_side = Side(border_style="thin", color=BORDER_COLOR)
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_bottom_double = Border(bottom=Side(border_style="double", color="203764"), top=thin_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # ==========================================
    # TAB 1: PREGLED IZVJEŠTAJA (DASHBOARD)
    # ==========================================
    ws1 = wb.create_sheet(title="Pregled Izvještaja")
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "FINANCIJSKI IZVJEŠTAJ - ANALIZA"
    ws1["A1"].font = font_title
    ws1["A2"] = "Automatski generirano iz baze podataka 'podaci.csv'"
    ws1["A2"].font = font_subtitle

    ws1["A4"] = "Metrike analizirane imovine"
    ws1["A4"].font = font_section

    headers_kpi = ["Metrika / Ticker"] + tickers
    for col_idx, text in enumerate(headers_kpi, start=1):
        cell = ws1.cell(row=5, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    pocetne = [df_prices[t].iloc[0] for t in tickers]
    zadnje = [df_prices[t].iloc[-1] for t in tickers]
    ukupni_povrati = [((df_prices[t].iloc[-1] / df_prices[t].iloc[0]) - 1) for t in tickers]
    prosjecni_dnevni = [df_returns[t].mean() for t in tickers]
    dnevna_vol = [df_returns[t].std() for t in tickers]
    anualizirana_vol = [df_returns[t].std() * np.sqrt(252) for t in tickers]

    kpi_data = [
        ("Početna cijena", pocetne, "$#,##0.00"),
        ("Zadnja cijena", zadnje, "$#,##0.00"),
        ("Ukupni povrat (%)", ukupni_povrati, "0.00%"),
        ("Prosječni dnevni prinos", prosjecni_dnevni, "0.000%"),
        ("Dnevna volatilnost (Std Dev)", dnevna_vol, "0.000%"),
        ("Anualizirana volatilnost", anualizirana_vol, "0.00%"),
    ]

    for row_idx, (metric, values, num_format) in enumerate(kpi_data, start=6):
        cell_metric = ws1.cell(row=row_idx, column=1, value=metric)
        cell_metric.font = font_bold if "Ukupni" in metric else font_regular
        cell_metric.border = border_all
        if row_idx % 2 == 1:
            cell_metric.fill = fill_zebra
            
        for val_idx, val in enumerate(values, start=2):
            cell_val = ws1.cell(row=row_idx, column=val_idx, value=val)
            cell_val.font = font_bold if "Ukupni" in metric else font_regular
            cell_val.number_format = num_format
            cell_val.alignment = align_right
            cell_val.border = border_all
            if row_idx % 2 == 1:
                cell_val.fill = fill_zebra

    for col_idx in range(1, len(tickers) + 2):
        ws1.cell(row=12, column=col_idx).border = border_bottom_double

    ws1["A14"] = "Struktura i opis"
    ws1["A14"].font = font_section

    desc_text = [
        "Ovaj Excel izvještaj generiran je potpuno automatski.",
        "Sustav povlači podatke s Yahoo Finance API-ja, strukturira ih i zapisuje u 'podaci.csv'.",
        "Nakon toga, ovaj modul pretvara sirove podatke u poslovni izvještaj.",
        "",
        "Struktura radnih listova:",
        "  1. Pregled izvještaja - Ključne metrike i performanse.",
        "  2. Povijesni podatci - Kompletna baza povijesnih cijena s prinosima.",
        "  3. Statistika i korelacija - Matrica korelacije i deskriptivna statistika.",
        "  4. Tehnički Indikatori - SMA (20/50/100/200), RSI, MACD, Bollinger Bands.",
        "  5. Osnovne Informacije - Profil imovine i fundamentalni podatci s Yahoo Finance-a.",
        "  6. Vizualizacija - Ugrađeni grafički prikazi."
    ]

    for idx, line in enumerate(desc_text, start=15):
        ws1.cell(row=idx, column=1, value=line).font = font_regular

    # ==========================================
    # TAB 2: POVIJESNI PODACI
    # ==========================================
    ws2 = wb.create_sheet(title="Povijesni Podaci")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "Povijesne Cijene i dnevni prinosi"
    ws2["A1"].font = font_title

    headers_data = ["Datum"]
    for t in tickers:
        headers_data.extend([f"{t} Price", f"{t} Return"])

    for col_idx, text in enumerate(headers_data, start=1):
        cell = ws2.cell(row=3, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    for row_idx, date in enumerate(dates, start=4):
        c_date = ws2.cell(row=row_idx, column=1, value=date.strftime("%Y-%m-%d"))
        c_date.alignment = align_center
        c_date.border = border_all
        
        col_counter = 2
        for t_idx, t in enumerate(tickers):
            price = df_prices.loc[date, t]
            c_price = ws2.cell(row=row_idx, column=col_counter, value=price)
            c_price.number_format = "$#,##0.00"
            c_price.border = border_all
            
            c_ret = ws2.cell(row=row_idx, column=col_counter + 1)
            if row_idx == 4:
                c_ret.value = "-"
                c_ret.alignment = align_center
            else:
                prev_row = row_idx - 1
                p_letter = get_column_letter(col_counter)
                c_ret.value = f"=({p_letter}{row_idx}-{p_letter}{prev_row})/{p_letter}{prev_row}"
                c_ret.number_format = "0.00%"
                c_ret.alignment = align_right
            c_ret.border = border_all
            
            if row_idx % 2 == 1:
                c_price.fill = fill_zebra
                c_ret.fill = fill_zebra
                
            col_counter += 2

    # ==========================================
    # TAB 3: STATISTIKA I KORELACIJA
    # ==========================================
    ws3 = wb.create_sheet(title="Statistika i Korelacija")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = "Statistička analiza i matrica korelacije"
    ws3["A1"].font = font_title

    ws3["A3"] = "Korelacija dnevnih prinosa"
    ws3["A3"].font = font_section

    corr_headers = ["Ticker"] + tickers
    for col_idx, text in enumerate(corr_headers, start=1):
        cell = ws3.cell(row=4, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    corr_matrix = df_returns.corr()
    for r_idx, t_row in enumerate(tickers, start=5):
        cell_r = ws3.cell(row=r_idx, column=1, value=t_row)
        cell_r.font = font_bold
        cell_r.border = border_all
        for c_idx, t_col in enumerate(tickers, start=2):
            corr_val = corr_matrix.loc[t_row, t_col]
            cell_val = ws3.cell(row=r_idx, column=c_idx, value=corr_val)
            cell_val.number_format = "0.00"
            cell_val.border = border_all
            cell_val.alignment = align_right
            
            if t_row == t_col:
                cell_val.fill = fill_accent

    if len(tickers) > 0:
        zadnji_red_corr = 4 + len(tickers)
        zadnji_stupac_corr = 1 + len(tickers)
        zadnje_slovo_corr = get_column_letter(zadnji_stupac_corr)
        raspon_matrice = f"B5:{zadnje_slovo_corr}{zadnji_red_corr}"
        
        rule = ColorScaleRule(
            start_type='num', start_value=-1.0, start_color='F8CBAD',
            mid_type='num', mid_value=0.5, mid_color='FFF2CC',
            end_type='num', end_value=1.0, end_color='C6E0B4'
        )
        ws3.conditional_formatting.add(raspon_matrice, rule)

    ws3["A10"] = "Deskriptivna Statistika"
    ws3["A10"].font = font_section

    stats_headers = ["Metrika"] + tickers
    for col_idx, text in enumerate(stats_headers, start=1):
        cell = ws3.cell(row=11, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    last_row = len(dates) + 3
    stats_rows = [
        ("Broj opservacija", "COUNT", "0"),
        ("Maksimalna cijena", "MAX", "$#,##0.00"),
        ("Minimalna cijena", "MIN", "$#,##0.00"),
        ("Prosječni dnevni prinos", "AVERAGE", "0.00%"),
        ("Dnevna volatilnost (Std Dev)", "STDEV.S", "0.00%"),
    ]

    for r_idx, (metric, formula_name, num_fmt) in enumerate(stats_rows, start=12):
        cell_metric = ws3.cell(row=r_idx, column=1, value=metric)
        cell_metric.font = font_bold
        cell_metric.border = border_all
        if r_idx % 2 == 1:
            cell_metric.fill = fill_zebra
            
        for t_idx, t in enumerate(tickers):
            price_col_letter = get_column_letter(2 + t_idx * 2)
            return_col_letter = get_column_letter(3 + t_idx * 2)
            
            if formula_name in ["COUNT", "MAX", "MIN"]:
                formula = f"={formula_name}('Povijesni Podaci'!{price_col_letter}4:{price_col_letter}{last_row})"
            else:
                formula = f"={formula_name}('Povijesni Podaci'!{return_col_letter}5:{return_col_letter}{last_row})"
                
            cell_val = ws3.cell(row=r_idx, column=2 + t_idx, value=formula)
            cell_val.number_format = num_fmt
            cell_val.alignment = align_right
            cell_val.border = border_all
            if r_idx % 2 == 1:
                cell_val.fill = fill_zebra

    # ==========================================
    # TAB 4: TEHNIČKI INDIKATORI
    # ==========================================
    ws4 = wb.create_sheet(title="Tehnički Indikatori")
    ws4.views.sheetView[0].showGridLines = True

    ws4["A1"] = "Tehnička analiza i indikatori"
    ws4["A1"].font = font_title
    ws4["A2"] = "Pregled ključnih indikatora (SMA 20/50/100/200, RSI, MACD, Bollinger Bands)"
    ws4["A2"].font = font_subtitle

    headers_tech = ["Datum"]
    for t in tickers:
        headers_tech.extend([
            f"{t} Price", f"{t} SMA 20", f"{t} SMA 50", f"{t} SMA 100", f"{t} SMA 200",
            f"{t} RSI 14", f"{t} MACD", f"{t} MACD Signal", f"{t} BB Upper", f"{t} BB Lower"
        ])

    for col_idx, text in enumerate(headers_tech, start=1):
        cell = ws4.cell(row=3, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    tech_data = {}
    for t in tickers:
        close_series = df_prices[t]
        sma20 = close_series.rolling(window=20).mean()
        sma50 = close_series.rolling(window=50).mean()
        sma100 = close_series.rolling(window=100).mean()
        sma200 = close_series.rolling(window=200).mean()
        try:
            rsi14 = ta.rsi(close_series, length=14)
        except:
            rsi14 = pd.Series(np.nan, index=close_series.index)
        
        try:
            macd_df = ta.macd(close_series, fast=12, slow=26, signal=9)
            macd_val = macd_df.iloc[:, 0] if macd_df is not None and not macd_df.empty else pd.Series(np.nan, index=close_series.index)
            macd_sig = macd_df.iloc[:, 2] if macd_df is not None and not macd_df.empty else pd.Series(np.nan, index=close_series.index)
        except:
            macd_val = macd_sig = pd.Series(np.nan, index=close_series.index)
        
        try:
            bb_df = ta.bbands(close_series, length=20, std=2)
            bb_lower = bb_df.iloc[:, 0] if bb_df is not None and not bb_df.empty else pd.Series(np.nan, index=close_series.index)
            bb_upper = bb_df.iloc[:, 2] if bb_df is not None and not bb_df.empty else pd.Series(np.nan, index=close_series.index)
        except:
            bb_lower = bb_upper = pd.Series(np.nan, index=close_series.index)
        
        tech_data[t] = {
            'price': close_series, 'sma20': sma20, 'sma50': sma50, 'sma100': sma100, 'sma200': sma200,
            'rsi14': rsi14, 'macd': macd_val, 'macd_sig': macd_sig, 'bb_upper': bb_upper, 'bb_lower': bb_lower
        }

    for row_idx, date in enumerate(dates, start=4):
        c_date = ws4.cell(row=row_idx, column=1, value=date.strftime("%Y-%m-%d"))
        c_date.alignment = align_center
        c_date.border = border_all
        
        col_counter = 2
        for t in tickers:
            ti = tech_data[t]
            row_vals = [
                (ti['price'].loc[date], "$#,##0.00"), (ti['sma20'].loc[date], "$#,##0.00"),
                (ti['sma50'].loc[date], "$#,##0.00"), (ti['sma100'].loc[date], "$#,##0.00"),
                (ti['sma200'].loc[date], "$#,##0.00"), (ti['rsi14'].loc[date], "0.00"),
                (ti['macd'].loc[date], "0.0000"), (ti['macd_sig'].loc[date], "0.0000"),
                (ti['bb_upper'].loc[date], "$#,##0.00"), (ti['bb_lower'].loc[date], "$#,##0.00")
            ]
            for sub_idx, (val, num_fmt) in enumerate(row_vals):
                cell = ws4.cell(row=row_idx, column=col_counter + sub_idx)
                if pd.isna(val) or val is None:
                    cell.value = "-"
                    cell.alignment = align_center
                else:
                    cell.value = float(val)
                    cell.number_format = num_fmt
                    cell.alignment = align_right
                cell.border = border_all
                if row_idx % 2 == 1:
                    cell.fill = fill_zebra
            col_counter += 10

    # ==========================================
    # TAB 5: OSNOVNE INFORMACIJE (NOVI MODUL)
    # ==========================================
    ws_info = wb.create_sheet(title="Osnovne Informacije")
    ws_info.views.sheetView[0].showGridLines = True

    ws_info["A1"] = "OSNOVNE INFORMACIJE I PROFIL IMOVINE"
    ws_info["A1"].font = font_title
    ws_info["A2"] = "Podaci preuzeti s Yahoo Finance API-ja"
    ws_info["A2"].font = font_subtitle

    headers_info = [
        "Ticker", "Naziv", "Tip", "Sektor", "Industrija", 
        "Trenutna Cijena", "Tržišna Kap.", "52t Visoko", "52t Nisko", 
        "P/E Omjer", "Div. Prinos (%)", "Valuta", "Burza"
    ]

    for col_idx, text in enumerate(headers_info, start=1):
        cell = ws_info.cell(row=4, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    print("[PROCES] Dohvaćam fundamentalne info podatke za tickere...")
    for row_idx, t in enumerate(tickers, start=5):
        try:
            info = yf.Ticker(t).info
        except:
            info = {}

        naziv = info.get("longName", info.get("shortName", "N/A"))
        tip = info.get("quoteType", "N/A")
        sektor = info.get("sector", "N/A")
        industrija = info.get("industry", "N/A")
        cijena = info.get("currentPrice", info.get("regularMarketPrice", "N/A"))
        market_cap = info.get("marketCap", "N/A")
        high_52 = info.get("fiftyTwoWeekHigh", "N/A")
        low_52 = info.get("fiftyTwoWeekLow", "N/A")
        pe_ratio = info.get("trailingPE", "N/A")
        div_yield = info.get("dividendYield", 0)
        
        if isinstance(div_yield, (int, float)):
            div_yield = div_yield * 100
        else:
            div_yield = "N/A"
            
        valuta = info.get("currency", "N/A")
        burza = info.get("exchange", "N/A")

        row_vals = [
            (t, align_center, "@"),
            (naziv, align_left, "@"),
            (tip, align_center, "@"),
            (sektor, align_left, "@"),
            (industrija, align_left, "@"),
            (cijena, align_right, "$#,##0.00"),
            (market_cap, align_right, "$#,##0"),
            (high_52, align_right, "$#,##0.00"),
            (low_52, align_right, "$#,##0.00"),
            (pe_ratio, align_right, "0.00"),
            (div_yield, align_right, "0.00%"),
            (valuta, align_center, "@"),
            (burza, align_center, "@")
        ]

        for col_idx, (val, alignment, num_fmt) in enumerate(row_vals, start=1):
            cell = ws_info.cell(row=row_idx, column=col_idx, value=val if val != "N/A" else "-")
            cell.font = font_regular
            cell.alignment = alignment
            cell.border = border_all
            if isinstance(val, (int, float)):
                cell.number_format = num_fmt
            if row_idx % 2 == 1:
                cell.fill = fill_zebra

    # ==========================================
    # TAB 6: VIZUALIZACIJA (GRAFIKONI)
    # ==========================================
    ws6 = wb.create_sheet(title="Vizualizacija")
    ws6.views.sheetView[0].showGridLines = False

    ws6["B2"] = "VIZUALIZACIJA PODATAKA"
    ws6["B2"].font = Font(name="Arial", size=20, bold=True, color="203764")
    ws6["B3"] = "Generirani svi tehnički i statistički grafikoni"
    ws6["B3"].font = font_subtitle

    popis_slika = [
        "Svi_tickeri_Bollinger_kompletno.png", "Svi_tickeri_EMA.png", "Svi_tickeri_histogrami.png",
        "Korelacijska matrica cijene.png", "Kumulativni prikaz.png", "Svi_tickeri_MACD.png",
        "rolling_volatilnost_podgrafovi.png", "Svi_tickeri_RSI.png", "Svi_tickeri_Sezonalnost_dani.png",
        "Svi_tickeri_Sezonalnost_mjeseci.png", "Svi_tickeri_SMA.png"
    ]

    trenutni_red = 5
    stupci_za_slike = ["B", "L"]
    brojac_slika = 0

    print("[PROCES] Ugrađivanje grafikona u Excel galeriju...")
    for naziv_slike in popis_slika:
        if os.path.exists(naziv_slike):
            img = OpenpyxlImage(naziv_slike)
            img.width = 550
            img.height = 300
            
            pozicija_stupac = stupci_za_slike[brojac_slika % 2]
            celija_sidro = f"{pozicija_stupac}{trenutni_red}"
            
            opis_celija = ws6.cell(row=trenutni_red - 1, column=2 if pozicija_stupac == "B" else 12)
            opis_celija.value = f"Grafikon: {naziv_slike.replace('.png', '').replace('_', ' ').upper()}"
            opis_celija.font = font_bold
            
            ws6.add_image(img, celija_sidro)
            if brojac_slika % 2 == 1:
                trenutni_red += 20
            brojac_slika += 1
        else:
            print(f"  [-] Preskočeno: {naziv_slike} nije pronađen.")

    if brojac_slika == 0:
        ws6["B5"] = "Nema pronađenih grafikona (.png datoteka) u mapi projekta."

    # Automatsko podešavanje širine stupaca
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if val_str.startswith('='):
                    val_str = "Formula_Length_Est"
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(excel_path)
    print(f"[USPJEH] Excel izvještaj je spremljen pod: {excel_path}")

if __name__ == "__main__":
    generiraj_excel_izvjestaj()